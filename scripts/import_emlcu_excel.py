#!/usr/bin/env python3
# =============================================================================
# Copyright (c) 2026 ComplyNex (complynex.online) — All Rights Reserved.
#
# This software is the proprietary and confidential intellectual property of
# ComplyNex (complynex.online), owned by Mohamed Emam.
# Unauthorized copying, reproduction, modification,
# distribution, reverse engineering, or use of this software, in whole or in
# part, via any medium, is strictly prohibited without prior written
# permission from the owner.
#
# Owner and Author: Mohamed Emam — GRC & Compliance Expert
# Contact: f.mohemam85@gmail.com | admin@complynex.online
#
# Registered deployment: complynex.online
# ComplyNex Full Compliance Platform (CMS) — Proprietary & Confidential
# =============================================================================

"""
FIXED EMLCU Excel File Importer
"""
import os
import sys
import hashlib
import logging
import json as json_module
from datetime import datetime
import openpyxl
import jellyfish
from screening.complynex_engine import normalize_arabic_text

# Setup logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler('logs/emlcu_import.log', encoding='utf-8')
    ]
)
logger = logging.getLogger(__name__)

def normalize_arabic_name(name):
    if not name: return None
    return normalize_arabic_text(str(name))

def generate_source_id(name, sheet_type):
    normalized = (name or '').strip().lower()
    hash_val = hashlib.md5(f"{sheet_type}_{normalized}".encode('utf-8')).hexdigest()[:16]
    return f"emlcu_{sheet_type}_{hash_val}"

def parse_terrorists_sheet(ws):
    entities = []
    headers = {}
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    for idx, col_name in enumerate(header_row):
        if col_name:
            headers[str(col_name).strip()] = idx
    
    has_arabic_headers = 'الاسم' in headers
    start_row = 2
    
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if not row or not any(row): continue
        
        def get_val(col_name):
            idx = headers.get(col_name)
            if idx is not None and idx < len(row):
                val = row[idx]
                return str(val).strip() if val is not None else None
            return None

        if has_arabic_headers:
            name = normalize_arabic_name(get_val('الاسم'))
            alt_name = normalize_arabic_name(get_val('اسم الشهرة'))
            dob = get_val('تاريخ الميلاد')
            national_id = get_val('الرقم القومي')
            case_number = get_val('رقم القرار')
            passport = get_val('جواز السفر')
            nationality = get_val('الجنسية')
        else:
            if not row[0]: continue
            name = normalize_arabic_name(row[0])
            alt_name = normalize_arabic_name(row[1]) if len(row) > 1 else None
            national_id = str(row[2]) if len(row) > 2 else None
            passport = str(row[3]) if len(row) > 3 else None
            nationality = str(row[4]) if len(row) > 4 else None
            case_number = str(row[5]) if len(row) > 5 else None
            dob = None

        if not name: continue
        
        entity = {
            'full_name': name,
            'original_name': name,
            'aliases': [alt_name] if alt_name else [],
            'entity_type': 'PERSON',
            'list_type': 'SANCTIONS',
            'source_id': generate_source_id(name, 'person'),
            'risk_level': 'EXTREME',
            'properties': {
                'case_number': case_number,
                'date_of_birth': dob,
                'national_id': national_id,
                'passport': passport,
                'nationality': nationality
            }
        }
        entities.append(entity)
    return entities

def parse_entities_generic(ws, sheet_type):
    entities = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]: continue
        name = normalize_arabic_name(row[1])
        if not name: continue
        entity = {
            'full_name': name,
            'original_name': name,
            'aliases': [],
            'entity_type': 'ORGANIZATION',
            'list_type': 'SANCTIONS',
            'source_id': generate_source_id(name, sheet_type),
            'risk_level': 'EXTREME',
            'properties': {'sheet': sheet_type}
        }
        entities.append(entity)
    return entities

def import_to_database(conn, db_type, entities):
    cursor = conn.cursor()
    
    # Use 'emlcu' as the main source for compatibility
    target_source = 'emlcu'
    
    cursor.execute(f"SELECT COUNT(*) FROM screening_watchlist WHERE source = '{target_source}'")
    before_count = cursor.fetchone()[0]
    
    records = []
    for entity in entities:
        name = entity['full_name']
        name_normalized = normalize_arabic_text(name)
        name_upper = name_normalized.upper()
        soundex = jellyfish.soundex(name)
        metaphone = jellyfish.metaphone(name)
        name_tokens = json_module.dumps(name_normalized.split())
        
        records.append((
            entity['full_name'],
            entity['original_name'],
            json_module.dumps(entity['aliases']),
            entity['entity_type'],
            entity['list_type'],
            'emlcu_egypt', # dataset
            json_module.dumps(['terrorism', 'egypt']), # topics
            target_source, # source
            entity['source_id'],
            'https://www.mlcu.org.eg',
            name_normalized, # name_normalized
            soundex,
            metaphone,
            name_tokens,
            entity['risk_level'],
            json_module.dumps(entity['properties'])
        ))

    placeholder = '?' if db_type == 'sqlite' else '%s'
    
    query = f"""
        INSERT INTO screening_watchlist (
            full_name, original_name, aliases, entity_type,
            list_type, dataset, topics,
            source, source_id, source_url,
            name_normalized, name_soundex, name_metaphone, name_tokens,
            risk_level, properties
        ) VALUES (
            {", ".join([placeholder]*16)}
        )
    """
    
    if db_type == 'sqlite':
        # SQLite doesn't support ON CONFLICT in the same way for bulk sometimes, 
        # or we use INSERT OR REPLACE
        query = query.replace("INSERT INTO", "INSERT OR REPLACE INTO")
        cursor.executemany(query, records)
    else:
        cursor.executemany(query + " ON CONFLICT (source_id) DO UPDATE SET full_name = EXCLUDED.full_name", records)
        
    conn.commit()
    
    cursor.execute(f"SELECT COUNT(*) FROM screening_watchlist WHERE source = '{target_source}'")
    after_count = cursor.fetchone()[0]
    
    stats = {'added': after_count - before_count, 'total': len(records)}
    logger.info(f"Import results: {stats}")
    return stats

def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('file_path')
    args = parser.parse_args()
    
    file_path = args.file_path
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    logger.info(f"Loading workbook: {file_path}")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    all_entities = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if 'الإرهابيين' in sheet_name or 'الارهابيين' in sheet_name:
            all_entities.extend(parse_terrorists_sheet(ws))
        elif 'الكيانات' in sheet_name:
            all_entities.extend(parse_entities_generic(ws, 'entity'))
            
    logger.info(f"Parsed {len(all_entities)} entities")
    
    import sqlite3
    conn = sqlite3.connect('complynex_cms.db')
    try:
        # CLEANUP: Delete old PEP-polluted records first if they are in 'emlcu'
        # But only if we have new data to replace them
        if all_entities:
            logger.info("Cleaning up old 'emlcu' records...")
            conn.execute("DELETE FROM screening_watchlist WHERE source = 'emlcu'")
            conn.commit()
            
        import_to_database(conn, 'sqlite', all_entities)
        print("Import successful!")
    finally:
        conn.close()

if __name__ == "__main__":
    main()
