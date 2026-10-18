# =============================================================================
# Copyright (c) 2026 ComplyNex (complynex.online) — All Rights Reserved.
#
# This software is the proprietary and confidential intellectual property of
# ComplyNex (complynex.online), owned by Mohamed Emam.
# Unauthorized copying, reproduction, modification,
# distribution, reverse engineering, or use of this software, in whole or in
# part, via any medium, is strictly prohibited without prior written
# permission from the owner.
#
# Owner and Author: Mohamed Emam — GRC & Compliance Expert
# Contact: f.mohemam85@gmail.com | admin@complynex.online
#
# Registered deployment: complynex.online
# ComplyNex Full Compliance Platform (CMS) — Proprietary & Confidential
# =============================================================================

"""Excel export of the carbon disclosure (bilingual headers, one sheet per section)."""

from openpyxl import Workbook
from openpyxl.styles import Font

from .engine import SCOPE3_CATEGORIES


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(60, max(len(str(c.value or '')) for c in col) + 2))
    return ws


def workbook(inv):
    wb = Workbook()
    wb.remove(wb.active)
    p, t = inv['period'], inv['totals']
    _sheet(wb, 'Summary', ['Item / البند', 'tCO2e'], [
        ('Reporting period / فترة التقرير', f"{p['label']} ({p['start_date']} – {p['end_date']})"),
        ('Consolidation approach / أسلوب التجميع', inv['cfg'].get('consolidation_approach') or 'operational_control'),
        ('Scope 1 / النطاق 1', t['scope1']),
        ('Scope 2 location-based / النطاق 2 (الموقع)', t['scope2_location']),
        ('Scope 2 market-based / النطاق 2 (السوق)', t['scope2_market']),
        ('Scope 3 / النطاق 3', t['scope3']),
        ('Total gross (location-based) / الإجمالي', t['total_location_based']),
        ('Total gross (market-based) / الإجمالي (السوق)', t['total_market_based']),
        ('Offsets retired (reported separately, not netted) / التعويضات', inv['offsets']['offsets_retired_tco2e']),
        ('Financed emissions Cat. 15 / الانبعاثات الممولة', inv['financed']['financed_tco2e']),
        ('PCAF weighted data-quality score / درجة جودة البيانات', inv['financed']['weighted_score']),
        ('Inventory status / حالة الجرد', inv['workflow'].get('status')),
    ])
    _sheet(wb, 'Scope 3', ['Category', 'Name', 'tCO2e', 'PCAF score'],
           [(k, SCOPE3_CATEGORIES.get(k, ''), v['tco2e'], v.get('pcaf_score')) for k, v in sorted(inv['scope3'].items())])
    _sheet(wb, 'Intensity', ['Metric / المؤشر', 'Value'], sorted(inv['intensity'].items()))
    _sheet(wb, 'Activity data', ['ID', 'Scope', 'Category', 'Activity', 'Month', 'Site', 'Quantity', 'Unit', 'Factor kgCO2e/unit',
                                 'Factor source', 'Factor vintage', 'Boundary share', 'Basis', 'tCO2e', 'Reviewed'],
           [(e['id'], e['scope'], e['category'], e['activity'], str(e.get('activity_month') or ''), e.get('site_name'),
             float(e['quantity'] or 0), e['unit'], float(e['factor_kgco2e_used']) if e.get('factor_kgco2e_used') is not None else None,
             e.get('factor_source_used'), e.get('factor_vintage_used'),
             float(e['boundary_share']) if e.get('boundary_share') is not None else None, e.get('basis'), e['tco2e'],
             'yes' if e.get('reviewed_by') else 'no') for e in inv['entries']])
    _sheet(wb, 'Factor sources', ['Source / المصدر', 'Vintage', 'Entries'],
           [(f['source'], f['vintage'], f['entries']) for f in inv['factors_used']])
    _sheet(wb, 'Offsets', ['Project', 'Registry ref.', 'tCO2e', 'Purchased', 'Retired on'],
           [(o['project_name'], o.get('registry_reference'), float(o['quantity_tco2e'] or 0), str(o.get('purchase_date') or ''),
             str(o.get('retired_on') or '')) for o in inv['offset_purchases']])
    _sheet(wb, 'Data quality', ['Severity', 'Code', 'Message'],
           [(i['severity'], i['code'], i['message']) for i in inv['dq']['issues']])
    _sheet(wb, 'Sign-offs', ['Action', 'From', 'To', 'User', 'At', 'Note'],
           [(e['action'], e['from_status'], e['to_status'], e['user_name'], str(e['at']), e.get('note')) for e in inv['events']])
    return wb
