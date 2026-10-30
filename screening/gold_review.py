"""
Gold Set human review (Phase 0 follow-up): the MLRO checks the stratified review sample
(gold_v<N>.review.jsonl, ~2 % of the Gold Set) by hand in Excel instead of a dedicated screen.

  build_xlsx()   -> the review workbook (one row per pair: what the pair tests, the screened
                    name, the listed name, the expected answer, and a column for the decision)
  parse_upload() -> the decisions from the filled workbook, checked against the sample

Decisions are stored next to the Gold Set (gold_v<N>.review_decisions.json) and applied to the
Gold Set by a separate, reviewed step: nothing here changes the Gold Set or the quality gate.
"""
import io
import json
import os
from datetime import datetime, timezone

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOLD_DIR = os.path.join(BASE, 'screening', 'data', 'gold')
VERSION = 1

OK, WRONG, UNSURE = 'صح', 'غلط', 'مش متأكد'
DECISIONS = (OK, WRONG, UNSURE)
EXPECTED = {1: 'نفس الشخص / نفس الكيان (لازم النظام يطلّعه)', 0: 'شخص مختلف (لازم النظام ما يطلّعوش)'}

CAT_AR = {
    'ar_al_prefix': 'اسم عربي - "ال" التعريف زيادة أو ناقصة (السيد / سيد)',
    'ar_colloquial': 'اسم عربي - كتابة عامية (مصطفا / مصطفي)',
    'ar_extra_name': 'اسم عربي - اسم زيادة في آخر اسم العميل',
    'ar_glued': 'اسم عربي - أسماء ملزوقة (عبدالرحمن / عبد الرحمن)',
    'ar_missing_later': 'اسم عربي - اسم ناقص من النص (الجد مثلاً)',
    'ar_orthography': 'اسم عربي - همزات / ياء / تاء مربوطة',
    'ar_sound_letters': 'اسم عربي - حروف متشابهة في النطق (ز / ظ ، س / ص)',
    'ar_titles': 'اسم عربي - لقب قبل الاسم (المهندس / الدكتور / الحاج)',
    'ar_typo': 'اسم عربي - خطأ إملائي (حرفين مبدّلين)',
    'company_forms': 'شركة - الشكل القانوني (Company / LLC / ش.م.م)',
    'cross_arabic_query_latin_listed': 'العميل بالعربي والمدرج بالإنجليزي',
    'cross_egyptian_latin': 'العميل بالإنجليزي (كتابة مصرية) والمدرج بالعربي',
    'cross_official': 'اسم أجنبي مكتوب بالعربي والمدرج بالإنجليزي',
    'lat_el_split': 'اسم إنجليزي - Abdel / Abdul / El مفصولة أو ملزوقة',
    'lat_order': 'اسم إنجليزي - ترتيب الاسم مختلف (اسم العائلة الأول)',
    'lat_spelling': 'اسم إنجليزي - تهجئة مختلفة (Mohamed / Muhammad)',
    'neg_company_similar': 'شركتين مختلفتين بأسماء متشابهة',
    'neg_cross_script_sibling': 'أخ / قريب بالإنجليزي مقابل المدرج بالعربي',
    'neg_family_diff': 'نفس الاسم الأول والأب والجد - عائلة مختلفة',
    'neg_lookalike_first': 'اسم أول شبيه (محمود / محمد) - باقي الاسم واحد',
    'neg_real_distinct_ar': 'شخصين مختلفين فعلاً (أرقام قومية مختلفة)',
    'neg_real_distinct_lat': 'شخصين أجانب مختلفين (سنين ميلاد مختلفة)',
    'neg_sibling': 'أخ للمدرج - نفس الأب والعائلة واسم أول مختلف',
    'sup_company_synonyms': 'إضافي - شركة بمرادفات (Trade / Trading)',
    'sup_cross_non_arabic_script': 'إضافي - اسم أجنبي بكتابة عربية غير مألوفة',
    'sup_lat_franco': 'إضافي - فرانكو (3 = ع ، 7 = ح)',
    'sup_lat_maghreb': 'إضافي - كتابة مغاربية / فرنسية',
}
HEADERS = ['#', 'ID', 'إيه اللي بنختبره', 'الاسم اللي اتفحص', 'الاسم في القايمة', 'الإجابة المتوقعة',
           'ملاحظة النظام', 'قرارك: الإجابة المتوقعة صح؟', 'تعليقك']


def review_path(version=VERSION):
    return os.path.join(GOLD_DIR, f'gold_v{version}.review.jsonl')


def decisions_path(version=VERSION):
    return os.path.join(GOLD_DIR, f'gold_v{version}.review_decisions.json')


def load_sample(path=None):
    with open(path or review_path(), encoding='utf-8') as f:
        return [json.loads(line) for line in f if line.strip()]


def build_xlsx(path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    rows = load_sample(path)
    wb = Workbook()
    ws = wb.active
    ws.title = 'المراجعة'
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, r in enumerate(rows, 1):
        ws.append([i, r['id'], CAT_AR.get(r['cat'], r['cat']), r['query'], r['listed'],
                   EXPECTED.get(int(r['label']), str(r['label'])), r.get('note') or '', '', ''])
    n = len(rows) + 1
    dv = DataValidation(type='list', formula1=f'"{OK},{WRONG},{UNSURE}"', allow_blank=True,
                        showErrorMessage=True, errorTitle='قيمة غير مقبولة',
                        error=f'اختار من القايمة: {OK} / {WRONG} / {UNSURE}')
    ws.add_data_validation(dv)
    dv.add(f'H2:H{n}')
    yellow = PatternFill('solid', fgColor='FFF2CC')
    for row in ws.iter_rows(min_row=2, max_row=n):
        row[7].fill = yellow
        row[8].fill = yellow
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    for col, w in zip('ABCDEFGHI', (6, 18, 34, 34, 34, 30, 30, 18, 30)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions['B'].hidden = True
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:I{n}'

    g = wb.create_sheet('الشرح')
    g.sheet_view.rightToLeft = True
    for line in [
        'كل سطر = اسم اتفحص + اسم موجود في قايمة + الإجابة اللي النظام لازم يوصلها.',
        f'المطلوب في عمود "قرارك": "{OK}" لو الإجابة المتوقعة صح، "{WRONG}" لو غلط، "{UNSURE}" لو محتاج بيانات أكتر.',
        'مثال: لو السطر بيقول "نفس الشخص" والاسمين في الحقيقة لشخصين مختلفين -> اختار "غلط" واكتب السبب في التعليق.',
        'ما تغيّرش أي عمود غير العمودين الأصفر، وما تمسحش ولا تعيد ترتيب السطور (عمود ID المخفي هو اللي بيربط قرارك بالسطر).',
        'مش لازم تخلص كله مرة واحدة: ارفع الملف في أي وقت، والسطور الفاضية بتتحسب "لسه ما اتراجعتش".',
        'الملف فيه أسماء من قوايم الحظر: يفضل يتشال من جهازك بعد الرفع.',
    ]:
        g.append([line])
    g.column_dimensions['A'].width = 120
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out, len(rows)


def parse_upload(fileobj, path=None):
    """(decisions {id: {'decision', 'note'}}, problems) from the filled workbook."""
    from openpyxl import load_workbook
    sample = {r['id']: r for r in load_sample(path)}
    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception as ex:
        return {}, [f'الملف مش Excel صالح ({ex.__class__.__name__})']
    ws = wb['المراجعة'] if 'المراجعة' in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or '').strip() for h in next(rows, [])]
    if head[:len(HEADERS)] != HEADERS:
        return {}, ['أعمدة الملف اتغيرت - نزّل نسخة جديدة من النظام وانقل قراراتك عليها']
    decisions, problems = {}, []
    for n, r in enumerate(rows, 2):
        if not r or not r[1]:
            continue
        rid = str(r[1]).strip()
        if rid not in sample:
            problems.append(f'سطر {n}: ID مش موجود في العينة')
            continue
        dec = str(r[7] or '').strip()
        note = str(r[8] or '').strip()
        if not dec and not note:
            continue
        if dec not in DECISIONS:
            problems.append(f'سطر {n}: القرار لازم يكون {OK} / {WRONG} / {UNSURE}')
            continue
        if dec == WRONG and not note:
            problems.append(f'سطر {n}: "{WRONG}" محتاج سبب في التعليق')
        if str(r[3] or '') != sample[rid]['query'] or str(r[4] or '') != sample[rid]['listed']:
            problems.append(f'سطر {n}: الأسماء في السطر اتغيرت عن الأصل')
            continue
        decisions[rid] = {'decision': {OK: 'correct', WRONG: 'wrong', UNSURE: 'unsure'}[dec], 'note': note[:500]}
    return decisions, problems


def save(decisions, user_id, version=VERSION):
    """Merge into the stored decisions (a later upload wins per pair); returns the totals."""
    p = decisions_path(version)
    try:
        cur = json.load(open(p, encoding='utf-8'))
    except (OSError, ValueError):
        cur = {'version': version, 'decisions': {}}
    now = datetime.now(timezone.utc).isoformat()
    for rid, d in decisions.items():
        cur['decisions'][rid] = dict(d, by=user_id, at=now)
    cur['updated_at'] = now
    tmp = p + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(cur, f, ensure_ascii=False, indent=1)
    os.replace(tmp, p)
    vals = [d['decision'] for d in cur['decisions'].values()]
    return {'reviewed': len(vals), 'correct': vals.count('correct'), 'wrong': vals.count('wrong'),
            'unsure': vals.count('unsure')}


def status(version=VERSION):
    try:
        total = sum(1 for _ in open(review_path(version), encoding='utf-8'))
    except OSError:
        return None
    try:
        vals = [d['decision'] for d in json.load(open(decisions_path(version), encoding='utf-8'))['decisions'].values()]
    except (OSError, ValueError, KeyError):
        vals = []
    return {'total': total, 'reviewed': len(vals), 'wrong': vals.count('wrong'), 'unsure': vals.count('unsure')}
