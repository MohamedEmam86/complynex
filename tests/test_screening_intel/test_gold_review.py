import io
import json

from screening import gold_review as GR

S = [{'id': 'a1', 'cat': 'ar_typo', 'label': 1, 'tier': 'core', 'query': 'سمر حجزاي', 'listed': 'سمر حجازي', 'note': ''},
     {'id': 'b2', 'cat': 'neg_sibling', 'label': 0, 'tier': 'core', 'query': 'سيد علي', 'listed': 'محمد علي', 'note': 'x'}]


def _sample(tmp_path):
    p = tmp_path / 'r.jsonl'
    p.write_text('\n'.join(json.dumps(x, ensure_ascii=False) for x in S), encoding='utf-8')
    return str(p)


def _fill(buf, values):
    from openpyxl import load_workbook
    wb = load_workbook(buf)
    ws = wb['المراجعة']
    for row, (dec, note) in values.items():
        ws.cell(row=row, column=8, value=dec)
        ws.cell(row=row, column=9, value=note)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def test_roundtrip(tmp_path):
    p = _sample(tmp_path)
    buf, n = GR.build_xlsx(p)
    assert n == 2
    d, probs = GR.parse_upload(_fill(buf, {2: (GR.OK, ''), 3: (GR.WRONG, 'نفس الشخص فعلاً')}), p)
    assert probs == []
    assert d == {'a1': {'decision': 'correct', 'note': ''}, 'b2': {'decision': 'wrong', 'note': 'نفس الشخص فعلاً'}}


def test_blank_rows_are_unreviewed_and_wrong_needs_reason(tmp_path):
    p = _sample(tmp_path)
    buf, _ = GR.build_xlsx(p)
    d, probs = GR.parse_upload(_fill(buf, {3: (GR.WRONG, '')}), p)
    assert 'a1' not in d and probs


def test_bad_value_rejected(tmp_path):
    p = _sample(tmp_path)
    buf, _ = GR.build_xlsx(p)
    d, probs = GR.parse_upload(_fill(buf, {2: ('yes', '')}), p)
    assert probs and 'a1' not in d


def test_not_excel(tmp_path):
    d, probs = GR.parse_upload(io.BytesIO(b'not a workbook'), _sample(tmp_path))
    assert d == {} and probs
